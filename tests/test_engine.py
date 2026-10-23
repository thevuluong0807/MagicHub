"""Chạy: .venv/bin/python tests/test_engine.py  (cần LibreOffice cho phần Office; không cần display)."""
import json
import os
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
os.environ["XDG_CONFIG_HOME"] = tempfile.mkdtemp()  # config riêng cho test, không đụng config thật

import pymupdf
from PIL import Image

from magichub.actions import available, by_id
from magichub.actions.base import parse_ranges
from magichub.geometry import R0, ring_hit


def run(aid, files, **params):
    return by_id(aid).run([Path(f) for f in files], params)


def make_pdf(path, pages=3, text="Hello Secret 123"):
    doc = pymupdf.open()
    for i in range(pages):
        pg = doc.new_page()
        pg.insert_text((72, 100), f"{text} page {i + 1}", fontsize=18)
    doc.save(str(path))
    return path


def pdf_text(path):
    d = pymupdf.open(str(path))
    return "".join(p.get_text() for p in d)


def test_geometry():
    assert ring_hit(0, 0, 3)[0] is None  # vùng chết
    assert ring_hit(0, -(R0 + 5), 3)[0] == 0  # 12h
    assert ring_hit(200, 100, 3)[0] == 1  # phải-dưới
    assert ring_hit(-200, 100, 3)[0] == 2  # trái-dưới
    assert ring_hit(-150, -80, 6)[0] == 5  # lệch trái 12h hơn 30°, 6 ô 60°
    assert ring_hit(0, -100, 0)[0] is None

def test_available(td):
    pdf, png = make_pdf(td / "a.pdf"), td / "a.png"
    Image.new("RGB", (80, 60), "red").save(png)
    ids = lambda files: {a.id for _, _, acts in available(files) for a in acts}
    assert "convert.pdf.docx" in ids([pdf]) and "edit.redact" in ids([pdf])
    assert "organize.merge" not in ids([pdf]) and "organize.merge" in ids([pdf, png])
    assert "convert.img.png" not in ids([png]) and "convert.img.jpg" in ids([png])
    assert "edit.redact" not in ids([png]) and "organize.zip" in ids([png])
    assert ids([td / "x.unknown"]) == {"organize.zip"}


def test_organize(td):
    pdf = make_pdf(td / "doc.pdf")
    (o,) = run("organize.split", [pdf], ranges="1-2,3")
    assert sorted(p.name for p in o.iterdir()) == ["doc_p1-2.pdf", "doc_p3.pdf"]
    (o,) = run("organize.pages", [pdf], ranges="2")
    assert len(pymupdf.open(str(o))) == 1 and "page 2" in pdf_text(o)
    (o,) = run("organize.rotate_cw", [pdf])
    assert pymupdf.open(str(o))[0].rotation == 90
    (o,) = run("organize.compress", [pdf])
    assert o.exists()
    other = make_pdf(td / "b.pdf", pages=2)
    (o,) = run("organize.merge", [pdf, other])
    assert len(pymupdf.open(str(o))) == 5
    (o,) = run("organize.zip", [pdf, other])
    assert o.suffix == ".zip"
    png = td / "i.png"
    Image.new("RGB", (100, 50), "blue").save(png)
    (o,) = run("organize.rotate_cw", [png])
    assert Image.open(o).size == (50, 100)
    (o,) = run("organize.merge", [pdf, png])
    assert len(pymupdf.open(str(o))) == 4


def test_compress_office_target(td):
    """docx/pptx/xlsx: nén ảnh trong media/, file mở lại được; target_mb tự tăng mức nén, không đạt thì báo lỗi và không để lại file."""
    import docx
    noise = Image.effect_noise((1800, 1400), 90).convert("RGB")
    jpg = td / "n.jpg"
    noise.save(jpg, quality=98)
    d = docx.Document()
    d.add_picture(str(jpg))
    src = td / "big.docx"
    d.save(src)
    (o,) = run("organize.compress", [src], level="Mạnh")
    assert o.stat().st_size < src.stat().st_size / 2 and len(docx.Document(str(o)).inline_shapes) == 1
    # target_mb: bắt đầu từ mức Nhẹ (còn quá lớn) và tự tăng mức tới khi đạt
    limit = (o.stat().st_size // 1048576) + 1
    (t,) = run("organize.compress", [src], level="Nhẹ (giữ nét)", target_mb=limit)
    assert t.stat().st_size <= limit * 1048576
    big = td / "pic.jpg"
    Image.effect_noise((4000, 3000), 90).convert("RGB").save(big, quality=100)
    assert big.stat().st_size > 3 * 1048576
    (u,) = run("organize.compress", [big], level="Nhẹ (giữ nét)", target_mb=2)
    assert u.stat().st_size <= 2 * 1048576
    from magichub.actions import organize
    real = organize._shrink
    organize._shrink = lambda f, out, level: out.write_bytes(b"x" * (3 * 1048576))  # nén kiểu gì cũng 3 MB
    try:
        run("organize.compress", [src], level="Cực mạnh", target_mb=1)
        raise AssertionError("phải báo lỗi")
    except ValueError as e:
        assert "3.0 MB" in str(e) and not list(td.glob("big_small*.tmp")) and len(list(td.glob("big_small*.docx"))) == 2  # 2 file thành công ở trên, không có file dở
    finally:
        organize._shrink = real


def test_anonymize_bad_zip(td):
    from magichub import anonymize as an
    f = Path(td) / "old.docx"
    f.write_bytes(b"\xd0\xcf\x11\xe0" + b"x" * 20)
    try:
        an.read_text(f)
        assert False
    except ValueError as e:
        assert "old.docx" in str(e) and "Office cũ" in str(e)


def test_detect_lowercase_path_and_military_org(td):
    from magichub import anonymize as an
    paras = ["Phần mềm BV3C", "Bảng bv3c.system", "BV3C hoạt động", "Báo cáo của Cục tác chiến điện tử.",
             "Tổng Công ty Công nghiệp Công nghệ cao Viettel", "Phòng ngừa rủi ro. Ban hành quy định."]
    got = {c.value: c for c in an.detect(paras)}
    assert got["BV3C"].kind == "PROJECT" and got["BV3C"].checked  # dạng thường trong tên bảng không làm mất mã
    assert got["Cục tác chiến điện tử"].kind == "ORG"
    assert got["Viettel"].kind == "ORG"
    assert "Phòng ngừa" not in got and "Ban hành" not in got


def test_word_lock(td):
    import base64
    import zipfile

    from magichub import wordlock

    def make(name, settings):
        f = Path(td) / name
        with zipfile.ZipFile(f, "w") as z:
            z.writestr("[Content_Types].xml", '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>')
            z.writestr("word/document.xml", "<d/>")
            z.writestr("word/settings.xml", settings)
            z.writestr("word/media/x.bin", bytes(range(256)) * 40)  # msoffcrypto cần file đủ lớn để giải mã lại được
        return f

    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    # 1) mật khẩu MỞ file: Protect -> container OLE; Unlock sai/đúng mật khẩu
    src = make("w.docx", f'<w:settings xmlns:w="{W}"/>')
    (enc,) = by_id("edit.protect").run([src], {"password": "pw123"})
    assert enc.read_bytes()[:4] == b"\xd0\xcf\x11\xe0"
    try:
        by_id("edit.unlock").run([enc], {"password": "x"})
        assert False
    except ValueError as e:
        assert "sai mật khẩu" in str(e)
    (d,) = by_id("edit.unlock").run([enc], {"password": "pw123"})
    assert zipfile.ZipFile(d).read("word/document.xml") == b"<d/>"
    # 2) khoá chỉ-đọc (Restrict Editing) có mật khẩu: gỡ đúng mật khẩu mới được
    salt = b"0123456789abcdef"
    h = base64.b64encode(wordlock._hash("s3cr3t", salt, "sha512", 100)).decode()
    ro = make("r.docx", f'<w:settings xmlns:w="{W}"><w:documentProtection w:edit="readOnly" w:enforcement="1" w:cryptAlgorithmSid="14" w:cryptSpinCount="100" '
                        f'w:hash="{h}" w:salt="{base64.b64encode(salt).decode()}"/></w:settings>')
    try:
        by_id("edit.unlock").run([ro], {"password": "wrong"})
        assert False
    except ValueError as e:
        assert "sai mật khẩu" in str(e)
    (u,) = by_id("edit.unlock").run([ro], {"password": "s3cr3t"})
    assert "documentProtection" not in zipfile.ZipFile(u).read("word/settings.xml").decode()
    assert by_id("edit.protect").accepts([src]) and not by_id("edit.protect").accepts([Path(td) / "a.doc"])


def test_anonymize(td):
    """Dò + mã hoá + Restore vòng tròn trên docx/xlsx/pptx/txt: không còn giá trị gốc, khôi phục đúng từng đoạn, giữ định dạng."""
    import docx, openpyxl, pptx
    from magichub import anonymize as an
    text = ("Ông Nguyễn Văn An là Giám đốc Công ty TNHH Sao Mai. Công ty TNHH Sao Mai ký hợp đồng BVL-2026-01 với Tập đoàn Viettel. "
            "Liên hệ an.nguyen@saomai.vn hoặc 0912 345 678. Dự án BVL-2026-01 do Nguyễn Văn An phụ trách; Tập đoàn Viettel cử Trần Thị Bình.")
    cands = {c.value: c for c in an.detect([text])}
    for v, k in (("Nguyễn Văn An", "PERSON"), ("Trần Thị Bình", "PERSON"), ("BVL-2026-01", "PROJECT"), ("an.nguyen@saomai.vn", "EMAIL"),
                 ("Công ty TNHH Sao Mai", "ORG"), ("Tập đoàn Viettel", "ORG")):
        assert v in cands and cands[v].kind == k and cands[v].checked, (v, list(cands))
    assert any(c.kind == "PHONE" for c in cands.values()) and cands["Nguyễn Văn An"].count == 2 and cands["BVL-2026-01"].count == 2

    # docx: tên bị tách thành nhiều run, có định dạng đậm; thêm header, bảng
    d = docx.Document()
    para = d.add_paragraph()
    for chunk in ("Giám đốc ", "Nguyễn ", "Văn"):
        para.add_run(chunk)
    para.runs[1].bold = True
    para.add_run(" An ký, công ty Sao Mai.")
    d.add_paragraph(text)
    d.add_table(1, 1).cell(0, 0).text = "Viettel: Tập đoàn Viettel"
    d.sections[0].header.paragraphs[0].text = "BVL-2026-01"
    d.core_properties.author = "Nguyễn Văn An"
    src = td / "hd.docx"
    d.save(src)
    entries = [{"value": v, "kind": c.kind} for v, c in cands.items() if c.checked]
    (anon, mp) = run("edit.anonymize", [src], entries=entries, custom=["Sao Mai"])
    got = "\n".join(an.read_text(anon))
    for secret in ("Nguyễn Văn An", "Viettel", "BVL-2026-01", "saomai.vn", "Sao Mai"):
        assert secret not in got, secret
    assert "[PERSON_1]" in got and "[ORG_" in got and "[PROJECT_1]" in got
    assert docx.Document(str(anon)).core_properties.author == ""  # metadata tác giả bị xoá
    assert openpyxl.load_workbook(mp)["Mapping"]["A2"].value.startswith("[")
    assert any(r.bold for r in docx.Document(str(anon)).paragraphs[0].runs)  # định dạng đậm của run vẫn còn
    # Restore: tự tìm mapping cạnh file; kết quả giống hệt bản gốc theo từng đoạn
    (back,) = run("edit.restore", [anon], mapping=str(mp), auto=True)
    assert an.read_text(back) == an.read_text(src)
    # AI bỏ ngoặc vuông / đổi hoa thường: vẫn khôi phục được
    t = td / "ai.txt"
    t.write_text("Theo PERSON_1 và [org_1], mã project_1 đã ký.", "utf-8")
    (rt,) = run("edit.restore", [t], mapping=str(mp))
    assert "Nguyễn Văn An" in rt.read_text("utf-8") and "[" not in rt.read_text("utf-8")
    # bảng mapping sai -> lỗi rõ ràng, không để lại file
    other = td / "other.txt"
    other.write_text("không có mã nào", "utf-8")
    try:
        run("edit.restore", [other], mapping=str(mp))
        raise AssertionError("phải báo lỗi")
    except ValueError:
        assert not list(td.glob("other_restored*"))

    # khớp không phân biệt hoa thường/dấu/khoảng trắng; biến thể có mã riêng và Restore trả đúng nguyên văn từng chỗ
    v = td / "var.docx"
    dv = docx.Document()
    dv.add_paragraph("Tập đoàn Viettel; TẬP ĐOÀN VIETTEL; tập  đoàn viettel; Tap doan Viettel và VIETTEL.")
    dv.add_paragraph("Nguyễn Văn An, NGUYỄN VĂN AN, nguyen van an.")
    dv.save(v)
    ents = [{"value": "Tập đoàn Viettel", "kind": "ORG"}, {"value": "Viettel", "kind": "ORG"}, {"value": "Nguyễn Văn An", "kind": "PERSON"}]
    (a, m) = run("edit.anonymize", [v], entries=ents)
    got = "\n".join(an.read_text(a))
    assert "iettel" not in got.lower() and "nguyen" not in got.lower() and "nguyễn" not in got.lower() and "đoàn" not in got.lower(), got
    ws = openpyxl.load_workbook(m)["Mapping"]
    rows = {r[1]: r[0] for r in ws.iter_rows(min_row=2, values_only=True)}
    assert rows["Tập đoàn Viettel"] == "[ORG_1]" and rows["TẬP ĐOÀN VIETTEL"].startswith("[ORG_1v") and len(set(rows.values())) == len(rows)
    (back,) = run("edit.restore", [a], mapping=str(m))
    assert an.read_text(back) == an.read_text(v), (an.read_text(back), an.read_text(v))
    assert any("Số lần" in str(c.value) for c in ws[1]) and "ORG_1" in " ".join(an.read_text(a))
    # dạng gọi tắt dùng cùng họ mã (ORG_1 / ORG_1s) mà Restore vẫn trả đúng từng dạng
    al = td / "alias.txt"
    al.write_text("Công ty TNHH Sao Mai ký. Sao Mai giao hàng, Công ty TNHH Sao Mai ok.", "utf-8")
    cs = {c.value: c for c in an.detect(al.read_text("utf-8").splitlines())}
    assert cs["Sao Mai"].alias_of == "Công ty TNHH Sao Mai", cs
    (aa, am) = run("edit.anonymize", [al])
    assert aa.read_text("utf-8") == "[ORG_1] ký. [ORG_1s] giao hàng, [ORG_1] ok.", aa.read_text("utf-8")
    (ar,) = run("edit.restore", [aa], mapping=str(am))
    assert ar.read_text("utf-8") == al.read_text("utf-8")

    # từ điển theo dự án: luôn áp dụng dù app không dò ra
    from magichub import config
    cfg = config.load()
    cfg["anonymize_dicts"] = {"BVL": ["ORG: Sao Hỏa", "Mã bí mật X"]}
    config.save(cfg)
    try:
        t2 = td / "d.txt"
        t2.write_text("Gửi SAO HỎA mã bí mật x.", "utf-8")
        (a2, m2) = run("edit.anonymize", [t2], entries=[{"value": "Gửi", "kind": "TERM"}], dictionary="BVL")
        assert "sao" not in a2.read_text("utf-8").lower() and "bí mật" not in a2.read_text("utf-8").lower()
        (r2,) = run("edit.restore", [a2], mapping=str(m2))
        assert r2.read_text("utf-8") == "Gửi SAO HỎA mã bí mật x."
    finally:
        cfg.pop("anonymize_dicts", None)
        config.save(cfg)

    # xlsx / pptx / txt
    wb = openpyxl.Workbook()
    wb.active["A1"] = "Khách hàng: Tập đoàn Viettel"
    wb.active["B2"] = "BVL-2026-01"
    wb.save(td / "b.xlsx")
    pr = pptx.Presentation()
    sl = pr.slides.add_slide(pr.slide_layouts[5])
    sl.shapes.title.text = "Dự án BVL-2026-01 của Viettel"
    pr.save(td / "s.pptx")
    (td / "n.md").write_text("# BVL-2026-01\nLiên hệ: Nguyễn Văn An\n", "utf-8")
    ents = [{"value": "Viettel", "kind": "ORG"}, {"value": "BVL-2026-01", "kind": "PROJECT"}, {"value": "Nguyễn Văn An", "kind": "PERSON"}]
    for name in ("b.xlsx", "s.pptx", "n.md"):
        f = td / name
        a, m = run("edit.anonymize", [f], entries=ents)
        assert not any("Viettel" in x or "BVL-2026-01" in x for x in an.read_text(a)), name
        (r,) = run("edit.restore", [a], mapping=str(m))
        assert an.read_text(r) == an.read_text(f), name
        assert by_id("edit.anonymize").accepts([f])
    assert not by_id("edit.anonymize").accepts([make_pdf(td / "p.pdf")]) and not by_id("edit.anonymize").accepts([td / "old.doc"])


def test_no_console_window(td):
    """Windows: mọi tiến trình con phải mang CREATE_NO_WINDOW (không thì nháy cửa sổ PowerShell/console khi xử lý)."""
    import subprocess
    from magichub.actions import base
    seen = []
    real_popen, real_flags = subprocess.Popen, base.NO_WINDOW

    class Fake:
        returncode = 0
        pid = 1
        def __init__(self, cmd, **kw): seen.append(kw)
        def communicate(self, timeout=None): return b"", b""

    subprocess.Popen, base.NO_WINDOW = Fake, {"creationflags": 0x08000000}
    try:
        base.run_proc(["powershell", "-File", "x.ps1"])
    finally:
        subprocess.Popen, base.NO_WINDOW = real_popen, real_flags
    assert seen[0].get("creationflags") == 0x08000000
    import inspect
    from magichub.actions import office
    assert "Hidden" in inspect.getsource(office._office_run)  # PowerShell tự ẩn cửa sổ dù cờ nào


def test_anonymize_detect_quality(td):
    """Dò ít nhầm: nhãn biểu mẫu/cụm thông dụng không phải tên; BVL là mã dự án; địa danh tách riêng; nhãn loại là 'Tổ chức'."""
    from magichub import anonymize as an
    text = """HƯỚNG DẪN SỬ DỤNG 11.BVL HS
Ngày soạn: 01/01/2026. Người duyệt: Trần Thị Bình. Người soạn: Nguyễn Văn An (Công ty TNHH Sao Mai).
Phạm vi áp dụng. Trong phạm vi này, hệ thống BVL hỗ trợ. Ban hành kèm theo quyết định; ban hành lần đầu.
Hệ Thống BVL chạy tại Hà Nội và Hà Nội. Tập đoàn Viettel và Viettel. Màn Hình Chính; màn hình chính.
Hồ Sơ dự án; hồ sơ thầu. Công ty TNHH Sao Mai ký với Sao Mai. BVL BVL BVL. Công ty Cổ phần ABC Việt ký."""
    got = {c.value: c for c in an.detect(text.splitlines())}
    for noise in ("Ngày soạn", "Phạm vi", "Phạm vi áp dụng", "Ban hành", "Hệ Thống", "Màn Hình Chính", "Hồ Sơ", "Người duyệt", "Người soạn"):
        assert noise not in got, noise
    assert got["BVL"].kind == "PROJECT" and got["BVL"].checked and got["BVL"].count >= 4  # 3 ký tự vẫn đủ nghĩa nếu là mã lặp lại
    assert got["Hà Nội"].kind == "PLACE" and not got["Hà Nội"].checked
    assert got["Trần Thị Bình"].kind == "PERSON" and got["Công ty TNHH Sao Mai"].kind == "ORG"
    assert got["Sao Mai"].alias_of == "Công ty TNHH Sao Mai"
    assert "Tập đoàn Viettel" in got and "Tập đoàn Viettel và Viettel" not in got  # "và" không nối hai tổ chức thành một
    assert "Công ty Cổ phần ABC Việt" in got and an.KIND_LABEL["ORG"] == "Tổ chức"
    assert all(len(c.value) >= an.MIN_LEN.get(c.kind, 2) for c in got.values())

    # tài liệu kiểu thật: tiêu đề in hoa là chữ Việt (không phải viết tắt), số mục không phải IP, tên tổ chức dài có chữ thường, không vắt qua dòng
    doc = """HÌNH ẢNH HIỆN TẠI CỦA HỆ THỐNG. DIỄN GIẢI. CAO BAY. HÌNH HÌNH HIỆN HIỆN TẠI TẠI DIỄN DIỄN CAO CAO BAY BAY
5.8.9.3 Huỷ mệnh lệnh; 1.1.1.4; Máy chủ 192.168.1.5 và 10.0.0.7 cổng 8080
Để đáp ứng nhiệm vụ, Tổng Công ty Công nghiệp Công nghệ cao Viettel – Chi nhánh phía Nam ký. Tập đoàn Công nghiệp – Viễn thông Quân đội sản xuất thiết bị.
Hệ thống MHTH kết nối UAV và GNSS. MHTH, MHTH, UAV, UAV, GNSS, GNSS.
Trường
Số Serial: 1; Sđ"""
    g = {c.value: c for c in an.detect(doc.splitlines())}
    for noise in ("HÌNH", "HIỆN", "TẠI", "DIỄN", "CAO", "BAY"):
        assert noise not in g, noise
    assert g["MHTH"].kind == "PROJECT" and g["MHTH"].checked  # viết tắt chưa rõ: coi là mã, chọn sẵn
    assert all(g[v].kind == "TECH" and not g[v].checked for v in ("UAV", "GNSS"))  # thuật ngữ công khai
    assert "192.168.1.5" in g and "10.0.0.7" in g and "5.8.9.3" not in g and "1.1.1.4" not in g
    assert "Tổng Công ty Công nghiệp Công nghệ cao Viettel" in g and "Tập đoàn Công nghiệp – Viễn thông Quân đội" in g
    assert not any("\n" in v or "Trường" in v for v in g)
    # mã ngắn/có số: chỉ cần lặp 2 lần là đủ nghĩa; mã kèm gạch có số ở giữa
    sm = {c.value: c for c in an.detect(["Dự án BVL và BV3C. Hệ thống BVL chạy BV3C-01, BV3C. Mã A1, A1; 4G 4G"])}
    assert sm["BVL"].kind == sm["BV3C"].kind == "PROJECT" and sm["BVL"].checked and sm["BV3C"].checked and "BV3C-01" in sm
    assert "A1" not in sm and "4G" not in sm  # cần ≥ 2 chữ cái và bắt đầu bằng chữ
    assert an.is_vn_syllable("CHÍ") and an.is_vn_syllable("quân") and not an.is_vn_syllable("BVL") and not an.is_vn_syllable("UAV")


def test_detect_term_vs_code(td):
    """Phân biệt Thuật ngữ (FPGA, TDOA) / Mã (BVL, BV3C) / Biến-định danh (FFT_WINDOW_FLAT) và danh sách 'SAFE' trong từ điển."""
    from magichub import anonymize as an, config
    from magichub.actions.anon import entries_from
    doc = """Hệ thống dùng FPGA và TDOA để định vị; FPGA, TDOA ok. Tham số FFT_WINDOW_FLAT, FFT_WINDOW_FLAT và getUserName, getUserName.
Dự án BVL: mã BVL-GN-CNP, BVL-GN-CKT. BVL BVL. TDOA (Time Difference of Arrival) và Time Difference of Arrival (TDOA). BV3C BV3C. XYZW XYZW. RADAR-01 RADAR-02 RADAR RADAR."""
    g = {c.value: c for c in an.detect(doc.splitlines())}
    for v in ("FPGA", "TDOA"):
        assert g[v].kind == "TECH" and not g[v].checked, v
    for v in ("BVL", "BV3C", "XYZW", "BVL-GN-CNP"):
        assert g[v].kind == "PROJECT" and g[v].checked, v
    assert g["FFT_WINDOW_FLAT"].kind == "IDENT" and not g["FFT_WINDOW_FLAT"].checked and g["getUserName"].kind == "IDENT"
    assert "FFT" not in g  # FFT nằm trong định danh, không tách ra
    assert g["RADAR"].kind == "PROJECT"  # tiền tố của mã RADAR-01 thắng bảng thuật ngữ
    assert an.is_vn_syllable("CAO") and "FPGA" in an.TECH_TERMS
    # "định nghĩa tiếng Anh ngay trong văn bản" -> thuật ngữ kể cả khi không có trong bảng
    d2 = {c.value: c for c in an.detect(["Dùng QZXV (Quick Zone Xray Value) cho QZXV. QZXV lại. BVL (Bộ Vô Lăng) BVL BVL."])}
    assert d2["QZXV"].kind == "TECH" and d2["BVL"].kind == "PROJECT"  # định nghĩa tiếng Việt: giữ là mã
    # SAFE: không dò và không bao giờ mã hoá, dù người dùng tick tay hoặc nằm trong từ điển khác
    assert "TDOA" not in {c.value for c in an.detect(doc.splitlines(), {an.fold_key("tdoa")})}
    from magichub import config as cfgmod
    real = cfgmod.config_dir
    cfgmod.config_dir = lambda: td
    try:
        cfg = cfgmod.load()
        cfg["anonymize_dicts"] = {"A": ["SAFE: TDOA", "ORG: Viettel"], "B": ["ORG: tdoa", "PROJECT: BVL"]}
        cfgmod.save(cfg)
        assert an.load_safe() == {"tdoa"}
        got = entries_from({"entries": [{"value": "TDOA", "kind": "TECH"}, {"value": "BVL", "kind": "PROJECT"}], "dictionary": "B"}, [])
        assert [e[0] for e in got] == ["BVL", "BVL"] and not any(e[0].lower() == "tdoa" for e in got)
        src = td / "s.txt"
        src.write_text("TDOA và BVL, Viettel", "utf-8")
        (o, m) = run("edit.anonymize", [src], entries=[{"value": "TDOA", "kind": "TECH"}, {"value": "BVL", "kind": "PROJECT"}], dictionary="A")
        assert o.read_text("utf-8") == "TDOA và [PROJECT_1], [ORG_1]", o.read_text("utf-8")
    finally:
        cfgmod.config_dir = real


def test_dictionary_dialogs(td):
    """Hộp thoại quản lý từ điển (tạo, thêm dòng, nhập file, lưu config) + checkbox 'Dùng từ điển' trong hộp Mã hoá."""
    from PySide6.QtCore import QTimer
    from PySide6.QtWidgets import (QApplication, QCheckBox, QComboBox, QDialog, QDialogButtonBox, QFileDialog, QInputDialog, QListWidget,
                                   QPushButton, QTableWidget, QTableWidgetItem)
    from magichub import anonymize as an, config, dialogs
    from magichub.actions.anon import entries_from

    app = QApplication.instance() or QApplication([])
    real_dir, real_text, real_open = config.config_dir, QInputDialog.getText, QFileDialog.getOpenFileName
    config.config_dir = lambda: td
    top = lambda: next(w for w in app.topLevelWidgets() if isinstance(w, QDialog) and w.isVisible())
    btn = lambda dlg, text: next(b for b in dlg.findChildren(QPushButton) if b.text() == text)
    try:
        # 1) quản lý: tạo từ điển "BVL", thêm 1 dòng tay + nhập file có dòng trùng
        imp = td / "terms.txt"
        imp.write_text("ORG: Viettel\nPERSON: Nguyễn Văn An\nviettel\nMã bí mật X\n", "utf-8")
        QInputDialog.getText = staticmethod(lambda *a, **k: ("BVL", True))
        QFileDialog.getOpenFileName = staticmethod(lambda *a, **k: (str(imp), ""))

        def drive_manager():
            d = top()
            btn(d, "Mới…").click()
            btn(d, "Thêm dòng").click()
            t = d.findChild(QTableWidget)
            t.setItem(0, 1, QTableWidgetItem("Sao Hỏa"))
            t.cellWidget(0, 0).setCurrentIndex(dialogs.DICT_KINDS.index("ORG"))
            btn(d, "Nhập từ file…").click()  # "viettel" trùng "Viettel" (không phân biệt hoa thường) nên bị bỏ
            d.findChild(QDialogButtonBox).button(QDialogButtonBox.Ok).click()

        QTimer.singleShot(300, drive_manager)
        saved = dialogs.ask_dictionaries()
        assert saved == {"BVL": ["ORG: Sao Hỏa", "ORG: Viettel", "PERSON: Nguyễn Văn An", "TERM: Mã bí mật X"]}, saved
        assert config.load()["anonymize_dicts"] == saved

        # 2) hộp Mã hoá: tick "Dùng từ điển" -> trả tên từ điển; bỏ tick -> rỗng; engine đọc các dòng thành mục
        cands = an.detect(["Công ty TNHH Sao Mai ký. Công ty TNHH Sao Mai ok."])

        def drive_anon(use):
            def go():
                d = top()
                c = next(x for x in d.findChildren(QCheckBox) if x.text().startswith("Dùng từ điển"))
                assert c.isEnabled() and c.isChecked()  # có từ điển -> mặc định bật
                c.setChecked(use)
                d.findChild(QDialogButtonBox).button(QDialogButtonBox.Ok).click()
            QTimer.singleShot(300, go)

        drive_anon(True)
        r = dialogs.ask_anonymize(cands)
        assert r["dictionary"] == "BVL" and r["custom"] == []
        got = entries_from(r, [])
        assert ("Sao Hỏa", "ORG") in got and ("Mã bí mật X", "TERM") in got
        drive_anon(False)
        assert dialogs.ask_anonymize(cands)["dictionary"] == "" and config.load()["anonymize_use_dict"] is False

        # 3) chưa có từ điển nào: checkbox bị khoá, không lỗi
        cfg = config.load()
        cfg["anonymize_dicts"] = {}
        config.save(cfg)

        def go2():
            d = top()
            c = next(x for x in d.findChildren(QCheckBox) if x.text().startswith("Dùng từ điển"))
            assert not c.isEnabled() and not c.isChecked()
            d.findChild(QDialogButtonBox).button(QDialogButtonBox.Ok).click()

        QTimer.singleShot(300, go2)
        assert dialogs.ask_anonymize(cands)["dictionary"] == ""
    finally:
        config.config_dir, QInputDialog.getText, QFileDialog.getOpenFileName = real_dir, real_text, real_open


def test_clipboard_anonymize(td):
    """Menu khay: mã hoá clipboard -> dán cho AI -> giải mã clipboard, bằng mapping gần nhất lưu cục bộ."""
    from PySide6.QtWidgets import QApplication
    from magichub import config, dialogs, theme
    from magichub.app import Hub

    app = QApplication.instance() or QApplication([])
    theme.apply(app)
    hub = Hub(app, hotkey_enabled=False)
    real_dir, real_ask = config.config_dir, dialogs.ask_anonymize
    config.config_dir = lambda: td
    notes = []
    hub.notify = lambda title, text, error=False: notes.append((title, error))
    dialogs.ask_anonymize = lambda cands: {"entries": [{"value": c.value, "kind": c.kind, "alias_of": c.alias_of} for c in cands if c.checked], "custom": [], "dictionary": "", "scrub": True}
    try:
        original = "Ông Nguyễn Văn An (Công ty TNHH Sao Mai) gửi an@saomai.vn. Sao Mai cam kết."
        app.clipboard().setText(original)
        hub.anon_clipboard()
        masked = app.clipboard().text()
        assert "Nguyễn" not in masked and "saomai" not in masked and "[PERSON_1]" in masked and "[ORG_1s]" in masked, masked
        reply = masked.replace("gửi", "đã nhận từ")  # "câu trả lời của AI" vẫn mang mã
        app.clipboard().setText(reply)
        hub.restore_clipboard()
        assert app.clipboard().text() == original.replace("gửi", "đã nhận từ")
        assert (td / "clipboard_mapping.json").exists() and not notes[-1][1]
        app.clipboard().setText("không có mã")
        hub.restore_clipboard()
        assert notes[-1][1]  # báo lỗi, không đổi clipboard
        assert app.clipboard().text() == "không có mã"
        hub.clear_clip_mapping()
        assert not (td / "clipboard_mapping.json").exists()
    finally:
        config.config_dir, dialogs.ask_anonymize = real_dir, real_ask
        hub.tray.hide()


def test_options(td):
    big = Image.effect_noise((1600, 1200), 80).convert("RGB")
    jpg = td / "n.jpg"
    big.save(jpg, quality=95)
    (a,), (b,) = run("organize.compress", [jpg], level="Nhẹ (giữ nét)"), run("organize.compress", [jpg], level="Mạnh")
    assert b.stat().st_size < a.stat().st_size and max(Image.open(b).size) <= 1600
    # output_dir: kết quả vào thư mục đích, không đụng file gốc
    dest = td / "out"
    dest.mkdir()
    from magichub import config
    cfg = config.load()
    cfg["output_dir"] = str(dest)
    config.save(cfg)
    (o,) = run("organize.compress", [jpg])
    assert o.parent == dest and jpg.exists()
    cfg["output_dir"] = None
    config.save(cfg)
    (o,) = run("organize.compress", [jpg])
    assert o.parent == td
    # OCR chỉ hiện khi có tesseract
    import shutil
    pdf = make_pdf(td / "o.pdf")
    assert ("convert.pdf.ocr" in [a.id for _, _, acts in available([pdf]) for a in acts]) == (shutil.which("tesseract") is not None)


def test_v03(td):
    pdf = make_pdf(td / "v.pdf", pages=3)
    marks = [{"type": "ink", "page": 0, "pts": [[0.1, 0.1], [0.5, 0.3], [0.8, 0.2]], "color": [1, 0, 0], "w": 2},
             {"type": "note", "page": 1, "x": 0.2, "y": 0.4, "text": "Kiểm tra lại số liệu\ndòng hai"}]
    (o,) = run("edit.markup", [pdf], marks=marks)
    d = pymupdf.open(str(o))
    assert len(list(d[0].annots())) == 1 and "Kiểm tra" in d[1].get_text()
    try:
        run("edit.markup", [pdf], marks=[])
        assert False
    except ValueError:
        pass
    (o,) = run("organize.reorder", [pdf], order=[2, 0])
    t = pdf_text(o)
    assert len(pymupdf.open(str(o))) == 2 and t.index("page 3") < t.index("page 1") and "page 2" not in t
    (o,) = run("organize.reorder", [pdf], order=[1, 0], rotate={"0": 90})
    d = pymupdf.open(str(o))
    assert d[0].rotation == 0 and d[1].rotation == 90  # trang gốc 0 (giờ ở vị trí 2) bị xoay
    # redact vùng tay: gỡ text dưới vùng, không cần từ khoá
    W, H = pymupdf.open(str(pdf))[0].rect.width, pymupdf.open(str(pdf))[0].rect.height
    (o,) = run("edit.redact", [pdf], rects=[{"page": 0, "x0": 0, "y0": 0.05, "x1": 1, "y1": 0.2}])
    d = pymupdf.open(str(o))
    assert "Secret" not in d[0].get_text() and "Secret" in d[1].get_text()
    # GUI offscreen: canvas bắt vẽ/vùng đúng toạ độ
    from PySide6.QtCore import QEvent, QPointF, Qt
    from PySide6.QtGui import QMouseEvent
    from PySide6.QtWidgets import QApplication
    from magichub.viewer import PageCanvas
    app = QApplication.instance() or QApplication([])
    cv = PageCanvas(pymupdf.open(str(pdf)), "ink")
    cv.resize(700, 800)
    box = cv._box()
    ev = lambda t, x, y: QMouseEvent(t, QPointF(x, y), Qt.LeftButton, Qt.LeftButton, Qt.NoModifier)
    cx, cy = box.x() + box.width() / 2, box.y() + box.height() / 2
    cv.mousePressEvent(ev(QEvent.MouseButtonPress, cx, cy))
    cv.mouseMoveEvent(ev(QEvent.MouseMove, cx + 50, cy + 40))
    cv.mouseReleaseEvent(ev(QEvent.MouseButtonRelease, cx + 50, cy + 40))
    assert len(cv.marks) == 1 and abs(cv.marks[0]["pts"][0][0] - 0.5) < 0.01
    cv.undo()
    assert not cv.marks
    cv.tool = "rect"
    cv.mousePressEvent(ev(QEvent.MouseButtonPress, cx, cy))
    cv.mouseReleaseEvent(ev(QEvent.MouseButtonRelease, cx + 100, cy + 100))
    assert cv.marks[0]["type"] == "rect" and cv.marks[0]["x1"] > cv.marks[0]["x0"]
    cv.grab()  # vẽ không crash
    # mọi loại mark: bấm/kéo trên canvas -> mark đúng loại; vẽ lại không crash
    cv.marks.clear()
    cv.tool, cv.shape = "shape", "arrow"
    cv.mousePressEvent(ev(QEvent.MouseButtonPress, cx - 40, cy))
    cv.mouseReleaseEvent(ev(QEvent.MouseButtonRelease, cx + 40, cy + 30))
    assert cv.marks[-1]["kind"] == "arrow" and cv.marks[-1]["x1"] > cv.marks[-1]["x0"]
    cv.all_pages, cv.stroke = True, 3.0
    cv.add_border()
    assert cv.marks[-1] == {"type": "border", "all": True, "color": cv._rgb(), "w": 3.0}
    cv.marks.append({"type": "text", "all": True, "x": 0.5, "y": 0.5, "text": "WM", "size": 30, "color": [1, 0, 0], "rot": 45})
    cv.marks.append({"type": "shape", "page": 0, "kind": "ellipse", "x0": .1, "y0": .1, "x1": .3, "y1": .2, "color": [0, 0, 1], "w": 2, "fill": True})
    cv.grab().save(str(td / "canvas.png"))
    # preview và PDF dùng chung engine: marks của canvas áp được vào PDF
    (o,) = run("edit.markup", [pdf], marks=cv.marks)
    assert "WM" in pymupdf.open(str(o))[2].get_text()


def test_integrate(td):
    from magichub import integrate
    if sys.platform != "linux":
        return
    made = integrate.install(td)
    names = {p.name for p in made}
    assert {"magichub.desktop", "MagicHub", "magichub.nemo_action"} <= names
    assert (td / ".local/share/nautilus/scripts/MagicHub").stat().st_mode & 0o111
    assert " menu " in (td / ".local/share/nemo/actions/magichub.nemo_action").read_text()
    assert len(integrate.remove(td)) == len(made) and not (td / ".config/autostart/magichub.desktop").exists()


def test_v05(td):
    from magichub import config
    # migrate config cũ (một hotkey) sang danh sách
    (config.config_dir() / "config.json").write_text(json.dumps({"hotkey": {"modifiers": ["ctrl"], "key": "m"}}))
    cfg = config.load()
    assert cfg["hotkeys"] == [{"modifiers": ["ctrl"], "key": "m", "group": None}] and "hotkey" not in cfg
    # ẩn + sắp xếp
    pdf = make_pdf(td / "m.pdf")
    ids = lambda: {g: [a.id for a in acts] for g, _, acts in available([pdf])}
    base = ids()
    assert "edit.crop" in base["edit"] and base["edit"][0] != "edit.crop"
    cfg["menu"] = {"hidden": ["edit.protect"], "order": ["edit.crop"]}
    config.save(cfg)
    now = ids()
    assert "edit.protect" not in now["edit"] and now["edit"][0] == "edit.crop"
    assert len(now["edit"]) == len(base["edit"]) - 1
    cfg["menu"] = {"hidden": [a for a in base["organize"]], "order": []}
    config.save(cfg)
    assert "organize" not in ids()  # nhóm rỗng tự biến mất
    cfg["menu"] = {"hidden": [], "order": []}
    config.save(cfg)
    # hộp thoại cài đặt (offscreen): OK giữ nguyên dữ liệu
    from PySide6.QtCore import QTimer
    from PySide6.QtWidgets import QApplication
    from magichub.settings import ask_hotkeys, ask_menu
    from PySide6.QtWidgets import QDialog
    app = QApplication.instance() or QApplication([])
    accept_open = lambda: [w.accept() for w in app.topLevelWidgets() if isinstance(w, QDialog) and w.isVisible()]
    QTimer.singleShot(400, accept_open)
    got = ask_menu({"hidden": ["edit.crop"], "order": []})
    assert got and "edit.crop" in got["hidden"] and len(got["order"]) == len(__import__("magichub.actions.base", fromlist=["x"]).REGISTRY)
    QTimer.singleShot(400, accept_open)
    assert ask_hotkeys(cfg["hotkeys"]) == cfg["hotkeys"]


def test_sandbox(td):
    import subprocess
    from magichub.actions.base import sandbox_prefix
    pre = sandbox_prefix([td])
    if not pre:
        return  # không có bwrap: bỏ qua
    run_ = lambda code: subprocess.run([*pre, sys.executable, "-c", code], capture_output=True).returncode
    assert run_(f"open({str(td / 'ok.txt')!r}, 'w').write('x')") == 0 and (td / "ok.txt").exists()
    assert run_(f"open({str(Path.home() / 'magichub_sandbox_probe')!r}, 'w')") != 0  # ngoài vùng cho phép: chỉ đọc
    assert run_("import socket; socket.create_connection(('1.1.1.1', 80), 2)") != 0  # không mạng


def test_convert(td):
    pdf = make_pdf(td / "c.pdf")
    (o,) = run("convert.pdf.txt", [pdf])
    assert "Secret" in o.read_text("utf-8")
    (o,) = run("convert.pdf.png", [pdf])
    assert o.is_dir() and len(list(o.glob("*.png"))) == 3
    (o,) = run("convert.pdf.pptx", [pdf])
    assert o.stat().st_size > 1000
    (o,) = run("convert.pdf.docx", [pdf])
    assert o.stat().st_size > 1000
    png = td / "x.png"
    Image.new("RGBA", (40, 40), (255, 0, 0, 0)).save(png)
    (o,) = run("convert.img.jpg", [png])
    assert Image.open(o).mode == "RGB"
    (o,) = run("convert.img.pdf", [png])
    assert len(pymupdf.open(str(o))) == 1
    # Office qua LibreOffice
    (d,) = run("convert.pdf.docx", [pdf])
    (o,) = run("convert.doc.pdf", [d])
    assert "Secret" in pdf_text(o)
    (o,) = run("convert.doc.txt", [d])
    assert "Secret" in o.read_text("utf-8")


def test_pdf_table(td):
    doc = pymupdf.open()
    pg = doc.new_page()
    for r in range(4):
        for c in range(3):
            rect = pymupdf.Rect(72 + c * 80, 72 + r * 30, 152 + c * 80, 102 + r * 30)
            pg.draw_rect(rect)
            pg.insert_text((rect.x0 + 5, rect.y0 + 20), f"r{r}c{c}", fontsize=11)
    doc.save(str(td / "t.pdf"))
    (o,) = run("convert.pdf.xlsx", [td / "t.pdf"])
    import openpyxl

    ws = openpyxl.load_workbook(o).worksheets[0]
    assert ws["B2"].value == "r1c1", ws["B2"].value


def test_edit(td):
    pdf = make_pdf(td / "e.pdf")
    (o,) = run("edit.protect", [pdf], password="pw1")
    d = pymupdf.open(str(o))
    assert d.needs_pass and d.authenticate("pw1")
    (o,) = run("edit.redact", [pdf], keywords="Secret")
    assert "Secret" not in pdf_text(o)
    try:
        run("edit.redact", [pdf], keywords="zzzz")
        raise AssertionError("phải lỗi khi không có kết quả")
    except ValueError:
        pass
    (o,) = run("edit.highlight", [pdf], keywords="Hello", color="Vàng")
    assert any(True for _ in pymupdf.open(str(o))[0].annots())
    stamp = td / "stamp.png"
    Image.new("RGBA", (100, 50), (0, 128, 0, 255)).save(stamp)
    marks = [
        {"type": "text", "all": True, "x": 0.5, "y": 0.5, "text": "BẢN NHÁP", "size": 40, "color": [0.85, 0.1, 0.1], "rot": 45},
        {"type": "text", "page": 0, "x": 0.1, "y": 0.2, "text": "Dòng một\nDòng hai", "size": 16, "color": [0, 0, 0]},
        {"type": "image", "all": True, "x": 0.7, "y": 0.9, "w": 0.2, "path": str(stamp)},
        {"type": "shape", "page": 0, "kind": "rect", "x0": 0.1, "y0": 0.3, "x1": 0.4, "y1": 0.4, "color": [0, 0, 1], "w": 2, "fill": True},
        {"type": "shape", "page": 0, "kind": "ellipse", "x0": 0.5, "y0": 0.3, "x1": 0.8, "y1": 0.4, "color": [0, 0, 1], "w": 2},
        {"type": "shape", "page": 0, "kind": "arrow", "x0": 0.1, "y0": 0.5, "x1": 0.4, "y1": 0.6, "color": [1, 0, 0], "w": 2},
        {"type": "border", "all": True, "color": [0, 0, 0], "w": 3}]
    (o,) = run("edit.markup", [pdf], marks=marks)
    d = pymupdf.open(str(o))
    assert all("BẢN NHÁP" in pg.get_text() and len(pg.get_images()) == 1 for pg in d) and "Dòng hai" in d[0].get_text()
    assert "Dòng hai" not in d[1].get_text()
    assert len(d[0].get_drawings()) > len(d[1].get_drawings()) > 0
    d[0].get_pixmap(dpi=60).save(str(td / "markup.png"))
    (o,) = run("edit.crop", [pdf], box=(0.1, 0.1, 0.6, 0.5))
    r = pymupdf.open(str(o))[0].rect
    assert abs(r.width - 595 * 0.5) < 3 and abs(r.height - 842 * 0.4) < 3, r
    png = td / "p.png"
    Image.new("RGB", (200, 100), "white").save(png)
    (o,) = run("edit.crop", [png], box=(0, 0, 0.5, 0.5))
    assert Image.open(o).size == (100, 50)
    big = td / "big.png"
    Image.new("RGB", (400, 300), "white").save(big)
    (o,) = run("edit.markup", [big], marks=[
        {"type": "text", "page": 0, "x": 0.1, "y": 0.1, "text": "Hi", "size": 30, "color": [0, 0, 1]},
        {"type": "border", "page": 0, "color": [1, 0, 0], "w": 4}])
    im = Image.open(o).convert("RGB")
    assert im.size == (400, 300) and o.suffix == ".png"
    assert any(im.getpixel((200, y))[0] > 200 and im.getpixel((200, y))[1] < 60 for y in range(30))  # khung đỏ ở mép trên
    assert im.getpixel((200, 150)) == (255, 255, 255)


def test_v07(td):
    pdf = make_pdf(td / "a.pdf", pages=3)
    # header/footer: trang thường + trang xoay 90
    d = pymupdf.open(str(pdf)); d[1].set_rotation(90); d.saveIncr()
    (o,) = run("edit.headfoot", [pdf], header="HEAD", footer="{n} / {N}", align="Phải", size=10)
    d = pymupdf.open(str(o))
    assert "HEAD" in d[0].get_text() and "3 / 3" in d[2].get_text() and "2 / 3" in d[1].get_text()
    for pg in (d[0], d[1]):  # chữ phải nằm ngang trong hệ toạ độ hiển thị (trang xoay cũng vậy)
        ln = [l for b in pg.get_text("dict")["blocks"] for l in b.get("lines", []) if "/ 3" in "".join(s["text"] for s in l["spans"])][0]
        m = pg.rotation_matrix  # get_text dùng toạ độ gốc: đổi sang hệ hiển thị
        dv = pymupdf.Point(ln["dir"]) * m - pymupdf.Point(0, 0) * m
        bb = pymupdf.Rect(ln["bbox"]) * m
        assert abs(dv.x - 1) < 1e-3, dv
        assert bb.x0 > pg.rect.width * 0.7 and bb.y0 > pg.rect.height * 0.9, bb
    # flatten / strip / unlock
    d = pymupdf.open(str(pdf)); d.set_metadata({"author": "Bi Mat"}); d[0].add_text_annot((50, 50), "ghi chu"); d.save(str(td / "m.pdf"))
    (o,) = run("edit.flatten", [td / "m.pdf"])
    assert not list(pymupdf.open(str(o))[0].annots())
    (o,) = run("edit.strip", [td / "m.pdf"])
    assert not pymupdf.open(str(o)).metadata.get("author")
    (enc,) = run("edit.protect", [pdf], password="pw")
    try:
        run("edit.unlock", [enc], password="sai")
        raise AssertionError
    except ValueError:
        pass
    (o,) = run("edit.unlock", [enc], password="pw")
    assert not pymupdf.open(str(o)).needs_pass
    # compare: khác 1 trang
    b = make_pdf(td / "b.pdf", pages=3, text="Hello Secret 123")
    d = pymupdf.open(str(b)); d[1].insert_text((72, 300), "KHAC BIET", fontsize=30); d.save(str(td / "b2.pdf"))
    assert by_id("organize.compare").accepts([pdf, td / "b2.pdf"]) and not by_id("organize.compare").accepts([pdf])
    (o,) = run("organize.compare", [pdf, td / "b2.pdf"])
    px = [pymupdf.open(str(o))[i].get_pixmap(dpi=40) for i in range(3)]
    red = lambda pm: sum(1 for x in range(0, pm.width, 2) for y in range(0, pm.height, 2) if pm.pixel(x, y)[0] > 180 and pm.pixel(x, y)[1] < 80)
    assert red(px[0]) == 0 and red(px[1]) > 5, (red(px[0]), red(px[1]))
    # ảnh: resize, strip exif, crop giữ EXIF xoay, heic opener
    im = Image.new("RGB", (200, 100), "red")
    ex = Image.Exif(); ex[274] = 6; ex[0x010F] = "Cam"  # Orientation=6 -> xoay còn 100x200
    jpg = td / "ph.jpg"; im.save(jpg, exif=ex)
    (o,) = run("edit.resize", [jpg], mode="Phần trăm (%)", value=50)
    assert Image.open(o).size == (50, 100)
    (o,) = run("edit.resize", [jpg], mode="Chiều rộng (px)", value=40)
    assert Image.open(o).size == (40, 80)
    (o,) = run("edit.strip", [jpg])
    assert Image.open(o).size == (100, 200) and not dict(Image.open(o).getexif())
    (o,) = run("edit.crop", [jpg], box=(0, 0, 1, 0.5))
    assert Image.open(o).size == (100, 100)
    assert "heic" in __import__("magichub.actions.base", fromlist=["x"]).KINDS["img"]


def test_ask_form():
    """Mọi kiểu ô nhập (text/password/choice/int/multiline) phải đóng được bằng OK; lỗi trong slot accept() làm OK im lặng."""
    from PySide6.QtCore import QTimer
    from PySide6.QtWidgets import QApplication, QDialog, QDialogButtonBox, QLineEdit
    from magichub.dialogs import ask_form
    app = QApplication.instance() or QApplication([])

    def fill_ok(secret):
        def go():
            dlg = next(w for w in app.topLevelWidgets() if isinstance(w, QDialog) and w.isVisible())
            for e in dlg.findChildren(QLineEdit):
                e.setText(secret)
            dlg.findChild(QDialogButtonBox).button(QDialogButtonBox.Ok).click()
            QTimer.singleShot(300, lambda: dlg.reject() if dlg.isVisible() else None)  # kẹt thì thoát để test fail chứ không treo
        QTimer.singleShot(250, go)

    fill_ok("pw1")
    assert ask_form("Mật khẩu", [{"name": "password", "label": "MK", "type": "password", "required": True}]) == {"password": "pw1"}
    fill_ok("pw1")
    assert ask_form("Gỡ", [{"name": "password", "label": "MK", "type": "text", "required": True}]) == {"password": "pw1"}
    fill_ok("x")
    got = ask_form("Hỗn hợp", [{"name": "a", "label": "A", "type": "text"}, {"name": "c", "label": "C", "type": "choice", "choices": ["u", "v"]},
                               {"name": "n", "label": "N", "type": "int", "default": 3}, {"name": "m", "label": "M", "type": "multiline"}])
    assert got == {"a": "x", "c": "u", "n": 3, "m": ""}, got


def test_cropwidget():
    from PySide6.QtCore import QEvent, QPoint, QPointF, Qt
    from PySide6.QtGui import QMouseEvent, QPixmap
    from PySide6.QtWidgets import QApplication

    from magichub.dialogs import _CropLabel

    app = QApplication.instance() or QApplication([])
    from magichub.dialogs import _preview  # import trong hàm: lỗi đường dẫn import lộ ra ở đây
    pdf = make_pdf(Path(tempfile.mkdtemp()) / "c.pdf", pages=1)
    png = pdf.with_suffix(".png"); Image.new("RGB", (80, 40), "blue").save(png)
    assert _preview(pdf).width() > 100 and _preview(png).width() == 80
    lab = _CropLabel(QPixmap(400, 300))

    def ev(t, x, y):
        b = Qt.LeftButton if t != QEvent.MouseMove or lab.mode else Qt.NoButton
        e = QMouseEvent(t, QPointF(x, y), QPointF(x, y), Qt.LeftButton, b, Qt.NoModifier)
        {QEvent.MouseButtonPress: lab.mousePressEvent, QEvent.MouseMove: lab.mouseMoveEvent,
         QEvent.MouseButtonRelease: lab.mouseReleaseEvent}[t](e)

    def drag(a, b):
        ev(QEvent.MouseButtonPress, *a); ev(QEvent.MouseMove, *b); ev(QEvent.MouseButtonRelease, *b)

    f = lab.sel
    assert (f.left(), f.top(), f.right(), f.bottom()) == (0, 0, 400, 300)  # mặc định khung full
    drag((400, 300), (250, 200))  # kéo góc dưới-phải vào
    drag((0, 0), (50, 50))  # kéo góc trên-trái vào
    r = lab.sel
    assert (r.left(), r.top(), r.right(), r.bottom()) == (50, 50, 250, 200)
    drag((150, 120), (170, 130))  # kéo trong vùng = di chuyển
    assert (lab.sel.left(), lab.sel.top()) == (70, 60) and lab.sel.width() == 200
    drag((270, 130), (300, 130))  # kéo cạnh phải = đổi rộng
    assert lab.sel.right() == 300 and lab.sel.left() == 70 and lab.sel.height() == 150
    lab.set_ratio(1)
    assert abs(lab.sel.width() - lab.sel.height()) < 1e-6
    drag((10, 10), (500, 500))  # kéo ra ngoài ảnh bị kẹp trong ảnh, giữ tỉ lệ 1:1
    r = lab.sel
    assert r.right() <= 400 and r.bottom() <= 300 and abs(r.width() - r.height()) < 1e-6 and r.height() == 290


def test_progress_engine(td):
    from magichub.actions import base, office
    from magichub.actions.base import Cancelled, bind, report, scope, steps
    pdf = make_pdf(td / "p.pdf", pages=5)
    got = []
    bind(lambda f, t: got.append(f), None)
    run("organize.split", [pdf], ranges="")
    fr = [f for f in got if f is not None]
    assert fr == sorted(fr) and fr[-1] == 1.0 and len(fr) >= 6, fr
    got.clear()
    for i in steps(range(2)):  # phạm vi lồng nhau: 2 file x 2 trang -> 0, .25, .5, .75, 1
        for j in steps(range(2)):
            pass
    fr = [round(f, 2) for f in got if f is not None]
    assert fr == sorted(fr) and 0.25 in fr and 0.75 in fr and fr[-1] == 1.0, fr
    n = []
    bind(lambda f, t: n.append(1), lambda: len(n) >= 2)  # huỷ sau 2 lần báo
    try:
        run("organize.split", [pdf], ranges="")
        raise AssertionError("phải bị huỷ")
    except Cancelled:
        pass
    bind(None, None)
    # chọn engine: Linux không có Office -> auto dùng LibreOffice; ép office -> báo lỗi rõ
    from magichub import config
    calls = []
    orig = office.soffice_convert, office.available, office._office_run
    office.soffice_convert = lambda s, f, o: calls.append(("lo", f)) or Path(o) / "x"
    try:
        assert office.office_convert(td / "a.docx", "pdf", td) == td / "x" and calls == [("lo", "pdf")]
        office.available = lambda k: True
        office._office_run = lambda s, f, o: (_ for _ in ()).throw(RuntimeError("COM hỏng"))
        sys_plat = office.sys.platform
        office.sys.platform = "win32"  # giả lập: Office có nhưng lỗi -> auto rơi về LibreOffice
        calls.clear()
        office.office_convert(td / "a.docx", "pdf", td)
        assert calls == [("lo", "pdf")]
        cfg = config.load(); cfg["engine"] = "office"; config.save(cfg)
        try:
            office.office_convert(td / "a.docx", "pdf", td)
            raise AssertionError
        except RuntimeError as e:
            assert "COM" in str(e)
        cfg["engine"] = "libreoffice"; config.save(cfg)
        calls.clear()
        office.office_convert(td / "a.docx", "pdf", td)
        assert calls == [("lo", "pdf")]
        assert not office.use_office("img", "pdf") and not office.use_office("doc", "docx")
    finally:
        office.sys.platform = sys_plat
        office.soffice_convert, office.available, office._office_run = orig
        cfg = config.load(); cfg["engine"] = "auto"; config.save(cfg)


def test_updater_git(td):
    """Chạy từ clone git: apply() kéo commit mới (git pull --ff-only), mode() = 'git'."""
    import subprocess
    from magichub import updater
    g = lambda cwd, *a: subprocess.run(["git", "-c", "user.name=t", "-c", "user.email=t@t", *a], cwd=cwd, check=True, capture_output=True)
    src, dst = td / "src", td / "dst"
    src.mkdir()
    g(src, "init", "-q", "-b", "main")
    (src / "requirements.txt").write_text("")
    (src / "v.txt").write_text("1")
    g(src, "add", "-A"), g(src, "commit", "-qm", "1")
    g(td, "clone", "-q", str(src), "dst")
    (src / "v.txt").write_text("2")
    g(src, "commit", "-qam", "2")
    old = updater.ROOT
    updater.ROOT = dst
    try:
        assert updater.mode() == "git"
        assert updater.apply({"assets": {}}) == [dst] and (dst / "v.txt").read_text() == "2"
        (dst / "v.txt").write_text("sửa tay")  # sửa đổi cục bộ + commit mới xung đột: báo lỗi, không ghi đè
        (src / "v.txt").write_text("3")
        g(src, "commit", "-qam", "3")
        try:
            updater.apply({"assets": {}})
            raise AssertionError("phải báo lỗi")
        except RuntimeError:
            assert (dst / "v.txt").read_text() == "sửa tay"
    finally:
        updater.ROOT = old


def test_updater(td):
    import http.server, threading
    from magichub import updater
    from magichub.actions.base import Cancelled, bind

    assert updater.parse("v0.10.0") > updater.parse("0.9.3") and updater.parse("v0.8.1") == (0, 8, 1)
    rel = lambda tag: (lambda: {"tag_name": tag, "html_url": "http://x/r", "assets": [
        {"name": "MagicHub-x86_64.AppImage", "browser_download_url": "http://x/a"}]})
    assert updater.check("0.8.1", rel("v0.8.1")) is None and updater.check("0.8.1", rel("v0.8.0")) is None
    info = updater.check("0.8.1", rel("v0.9.0"))
    assert info["version"] == "0.9.0" and info["assets"]["MagicHub-x86_64.AppImage"] == "http://x/a"
    def boom(): raise OSError("offline")
    assert updater.check("0.8.1", boom) is None  # mất mạng: im lặng
    # cài AppImage từ server cục bộ: thay file đang chạy, có % tiến trình, không để lại .new/.part
    (td / "srv").mkdir()
    (td / "srv" / "new.AppImage").write_bytes(b"NEW" * 200000)
    h = lambda *a, **k: http.server.SimpleHTTPRequestHandler(*a, directory=str(td / "srv"), **k)
    srv = http.server.ThreadingHTTPServer(("127.0.0.1", 0), h)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    url = f"http://127.0.0.1:{srv.server_port}/new.AppImage"
    app = td / "MagicHub.AppImage"; app.write_bytes(b"OLD")
    os.environ["APPIMAGE"] = str(app)
    real = updater.sys.platform
    updater.sys.platform = "linux"
    try:
        fr = []
        bind(lambda f, t: fr.append(f), None)
        assert updater.mode() == "appimage"
        updater.apply({"assets": {"MagicHub-x86_64.AppImage": url}})
        assert app.read_bytes()[:3] == b"NEW" and os.access(app, os.X_OK) and fr[-1] == 1.0
        assert not list(td.glob("*.new")) and not list(td.glob("*.part"))
        app.write_bytes(b"OLD")  # huỷ giữa chừng: giữ nguyên bản cũ, dọn file tạm
        bind(None, lambda: True)
        try:
            updater.apply({"assets": {"MagicHub-x86_64.AppImage": url}})
            raise AssertionError
        except Cancelled:
            pass
        assert app.read_bytes() == b"OLD" and not list(td.glob("*.part")) and not list(td.glob("*.new"))
    finally:
        bind(None, None); updater.sys.platform = real; os.environ.pop("APPIMAGE"); srv.shutdown()


def test_docx_progress_hook():
    import logging
    from magichub.actions.base import Cancelled, bind
    from magichub.actions.convert import _PageLog

    got = []
    bind(lambda f, t: got.append(round(f, 2)), None)
    h = _PageLog()
    for i in (1, 2, 1, 2):  # 2 trang: pha phân tích rồi pha dựng
        h.emit(logging.LogRecord("x", logging.INFO, "", 0, f"({i}/2) Page {i}", None, None))
    assert got == [0.25, 0.5, 0.75, 1.0], got
    bind(None, lambda: True)
    try:
        h.emit(logging.LogRecord("x", logging.INFO, "", 0, "(1/2) Page 1", None, None))
        raise AssertionError
    except Cancelled:
        pass
    bind(None, None)


def test_hub_toast(td):
    import time as _t
    from PySide6.QtWidgets import QApplication
    from magichub import theme
    from magichub.actions.base import REGISTRY, Action, steps
    from magichub.app import Hub

    app = QApplication.instance() or QApplication([])
    theme.apply(app)
    hub = Hub(app, hotkey_enabled=False)

    def pump(cond, sec=5):
        end = _t.time() + sec
        while not cond() and _t.time() < end:
            app.processEvents()
        assert cond()

    def slow(files, p):
        for _ in steps(range(200)):
            _t.sleep(0.02)
        return [files[0]]

    act = Action("t.slow", "edit", "Slow", "chậm", slow)
    f = make_pdf(td / "s.pdf", 1)
    hub.files = [f]
    hub.on_chosen(act)
    pump(lambda: hub.toast.state == "running")  # sau 600 ms thẻ tiến trình hiện
    pump(lambda: hub.toast.frac is not None and hub.toast.frac > 0.1)
    assert "Slow" in hub.toast.title.text() and hub.toast.isVisible()
    hub.cancel_job(next(iter(hub.jobs)))
    pump(lambda: hub.toast.state == "error" and "huỷ" in hub.toast.title.text())
    assert not hub.jobs
    def mid(files, p):  # chạy lâu hơn 600 ms rồi xong bình thường: thẻ phải chuyển từ chạy sang xong (lỗi cũ: kẹt "đang chạy")
        for _ in steps(range(40)):
            _t.sleep(0.02)
        return []

    hub.on_chosen(Action("t.mid", "edit", "Mid", "vừa", mid))
    pump(lambda: hub.toast.state == "running")
    pump(lambda: hub.toast.state == "done" and "Mid" in hub.toast.title.text())
    assert not hub.jobs and hub.toast.close_btn.isVisible()

    def late(files, p):  # bước cuối không ngắt được: huỷ trước khi xong -> kết quả bị bỏ
        for _ in range(75):
            _t.sleep(0.02)
        o = td / "late_out.txt"
        o.write_text("x")
        return [o]

    hub.on_chosen(Action("t.late", "edit", "Late", "muộn", late))
    pump(lambda: hub.toast.state == "running")
    hub.cancel_job(next(iter(hub.jobs)))
    pump(lambda: hub.toast.state == "error" and not hub.jobs)
    assert not (td / "late_out.txt").exists()
    done = Action("t.fast", "edit", "Fast", "nhanh", lambda files, p: files)
    hub.on_chosen(done)
    pump(lambda: hub.toast.state == "done" and "Fast" in hub.toast.title.text())
    hub.toast.dismiss()
    hub.tray.hide()


if __name__ == "__main__":
    test_geometry()
    for t in (test_available, test_organize, test_options, test_v03, test_integrate, test_v05, test_sandbox, test_convert, test_pdf_table, test_edit, test_v07, test_compress_office_target, test_anonymize, test_anonymize_detect_quality, test_detect_term_vs_code, test_dictionary_dialogs, test_clipboard_anonymize, test_no_console_window, test_progress_engine, test_hub_toast, test_updater, test_updater_git):
        with tempfile.TemporaryDirectory() as d:
            t(Path(d))
        print("ok", t.__name__)
    test_docx_progress_hook()
    test_ask_form()
    test_cropwidget()
    print("ok test_cropwidget")
    print("ok test_geometry\nALL PASS")
