"""Nhóm Convert. Bảng đích theo loại file gốc ở TARGETS; chất lượng xem docs/CONSTRAINTS.md."""
from __future__ import annotations

import io
import shutil
from pathlib import Path

import pymupdf
from PIL import Image, ImageOps

from .base import (REGISTRY, Action, ext, natural_key, open_pdf, out_path,
                   soffice_convert, tmpdir, unique)

TARGETS = {
    "pdf": ["docx", "xlsx", "pptx", "png", "jpg", "txt"],
    "doc": ["pdf", "txt", "png", "jpg"],
    "xls": ["pdf", "csv", "png"],
    "ppt": ["pdf", "png", "jpg", "txt"],
    "img": ["png", "jpg", "webp", "pdf", "bmp"],
    "txt": ["pdf"],
}
KIND_NAME = {"pdf": "PDF", "doc": "Word", "xls": "Excel", "ppt": "PowerPoint", "img": "Ảnh", "txt": "Text"}
PIL_FMT = {"png": "PNG", "jpg": "JPEG", "webp": "WEBP", "bmp": "BMP"}


def render_pdf(pdf, to, base: Path) -> Path:
    """PDF -> ảnh. 1 trang: file; nhiều trang: thư mục page_001..."""
    doc = open_pdf(pdf)
    save = lambda page, p: page.get_pixmap(dpi=150, alpha=False).save(str(p))
    if len(doc) == 1:
        out = out_path(base, new_ext=to)
        save(doc[0], out)
    else:
        out = unique(base.with_name(f"{base.stem}_{to}"))
        out.mkdir()
        for i, page in enumerate(doc, 1):
            save(page, out / f"page_{i:03d}.{to}")
    doc.close()
    return out


def pdf_to_docx(src: Path) -> Path:
    from pdf2docx import Converter

    out = out_path(src, new_ext="docx")
    cv = Converter(str(src))
    try:
        cv.convert(str(out))
    finally:
        cv.close()
    return out


def pdf_to_xlsx(src: Path) -> Path:
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    doc = open_pdf(src)
    for pno, page in enumerate(doc, 1):
        for t, tab in enumerate(page.find_tables().tables, 1):
            ws = wb.create_sheet(f"P{pno}_T{t}")
            for row in tab.extract():
                ws.append([c or "" for c in row])
    if not wb.sheetnames:
        raise ValueError(f"{src.name}: không tìm thấy bảng nào (PDF scan chưa hỗ trợ)")
    out = out_path(src, new_ext="xlsx")
    wb.save(out)
    return out


def pdf_to_pptx(src: Path) -> Path:
    from pptx import Presentation

    doc = open_pdf(src)
    prs = Presentation()
    prs.slide_width, prs.slide_height = int(doc[0].rect.width * 12700), int(doc[0].rect.height * 12700)
    for page in doc:
        png = page.get_pixmap(dpi=150, alpha=False).tobytes("png")
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(io.BytesIO(png), 0, 0, prs.slide_width, prs.slide_height)
    out = out_path(src, new_ext="pptx")
    prs.save(out)
    return out


def pdf_to_txt(src: Path) -> Path:
    doc = open_pdf(src)
    out = out_path(src, new_ext="txt")
    out.write_text("\n\n".join(p.get_text() for p in doc), "utf-8")
    return out


def pptx_text(pptx: Path, out: Path) -> Path:
    from pptx import Presentation

    lines = []
    for n, slide in enumerate(Presentation(str(pptx)).slides, 1):
        lines.append(f"--- Slide {n} ---")
        lines += [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text]
    out.write_text("\n".join(lines), "utf-8")
    return out


def image_convert(src: Path, to: str) -> Path:
    im = ImageOps.exif_transpose(Image.open(src))
    if to in ("jpg", "bmp") and (im.mode in ("RGBA", "LA", "P")):
        im = im.convert("RGBA")
        bg = Image.new("RGB", im.size, "white")
        bg.paste(im, mask=im.getchannel("A"))
        im = bg
    elif im.mode not in ("RGB", "RGBA", "L"):
        im = im.convert("RGB")
    out = out_path(src, new_ext=to)
    im.save(out, PIL_FMT[to], **({"quality": 92} if to in ("jpg", "webp") else {}))
    return out


def images_to_pdf(files, out: Path) -> Path:
    doc = pymupdf.open()
    for f in files:
        im = ImageOps.exif_transpose(Image.open(f)).convert("RGB")
        buf = io.BytesIO()
        im.save(buf, "JPEG", quality=92)
        page = doc.new_page(width=im.width * 0.75, height=im.height * 0.75)  # 96 dpi -> pt
        page.insert_image(page.rect, stream=buf.getvalue())
    doc.save(str(out), deflate=True)
    doc.close()
    return out


def convert_one(src: Path, kind: str, to: str) -> Path:
    if kind == "pdf":
        return {"docx": pdf_to_docx, "xlsx": pdf_to_xlsx, "pptx": pdf_to_pptx,
                "txt": pdf_to_txt}.get(to, lambda s: render_pdf(s, to, s))(src)
    if kind == "img":
        return images_to_pdf([src], out_path(src, new_ext="pdf")) if to == "pdf" else image_convert(src, to)
    with tmpdir() as td:
        if to in ("png", "jpg"):
            return render_pdf(soffice_convert(src, "pdf", td), to, src)
        if kind == "ppt" and to == "txt":
            pptx = src if ext(src) == "pptx" else soffice_convert(src, "pptx", td)
            return pptx_text(pptx, out_path(src, new_ext="txt"))
        fmt = "txt:Text" if to == "txt" else to
        dest = out_path(src, new_ext=to)
        shutil.move(str(soffice_convert(src, fmt, td)), dest)
        return dest


def _register(kind: str, to: str):
    def run(files, params):
        return [convert_one(f, kind, to) for f in sorted(files, key=natural_key) if ext(f) != to]

    REGISTRY.append(Action(
        f"convert.{kind}.{to}", "convert", to.upper(), f"{KIND_NAME[kind]} → {to.upper()}", run,
        kinds=(kind,), check=lambda files: any(ext(f) != to for f in files)))


for _k, _targets in TARGETS.items():
    for _t in _targets:
        _register(_k, _t)
